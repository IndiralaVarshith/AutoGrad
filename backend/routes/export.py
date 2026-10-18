"""
AutoGrade — Export Routes

GET /export/{exam_id} — Generates Excel (.xlsx) report using openpyxl
Columns: Hall Ticket | Name | Q1 | Q2 ... | Total | Max Total | Flagged Qs | Avg Confidence | Status
"""

from __future__ import annotations

import io
from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from config import get_firestore


router = APIRouter(tags=["Export"])


# =========================================================
# GET /export/{exam_id}
# =========================================================

@router.get("/export/{exam_id}")
async def export_exam_results(exam_id: str):
    """
    Generate and stream an Excel (.xlsx) file containing all student results.
    Dynamically identifies all questions and builds tabular columns.
    """
    db = get_firestore()
    if db is None:
        raise HTTPException(status_code=503, detail="Firestore is unavailable.")

    try:
        docs = (
            db.collection("exams")
            .document(exam_id)
            .collection("results")
            .stream()
        )
        papers = [doc.to_dict() for doc in docs]
    except Exception as error:
        raise HTTPException(status_code=500, detail=f"Failed to fetch results: {error}")

    if not papers:
        raise HTTPException(status_code=404, detail=f"No results found for exam {exam_id}.")

    # Collect all unique question numbers across all papers
    question_nums = set()
    for paper in papers:
        for q in paper.get("questions", []):
            if "question_number" in q:
                question_nums.add(int(q["question_number"]))
    sorted_q_nums = sorted(list(question_nums))

    # Initialize workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Exam_{exam_id}"

    # Header styling
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB"),
    )

    # Build Header Row
    headers = ["Hall Ticket", "Name"]
    for q_num in sorted_q_nums:
        headers.append(f"Q{q_num}")
    headers.extend(["Total", "Max Total", "Flagged Qs", "Avg Confidence", "Status"])

    ws.append(headers)

    # Style Header
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    # Status fills
    status_fills = {
        "COMPLETE": PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid"),
        "NEEDS_REVIEW": PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"),
        "PROCESSING": PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"),
    }

    # Data Rows
    for row_idx, paper in enumerate(papers, start=2):
        hall_ticket = paper.get("hall_ticket", "")
        name = paper.get("student_name", "")
        total_score = paper.get("total_score", 0.0)
        max_total = paper.get("max_total", 0.0)
        flagged_count = paper.get("flagged_count", 0)
        avg_confidence = paper.get("avg_confidence", 0.0)
        status = paper.get("status", "COMPLETE")

        # Map scores per question
        q_map = {}
        for q in paper.get("questions", []):
            q_map[int(q.get("question_number", 0))] = q.get("final_score", 0.0)

        row_data = [hall_ticket, name]
        for q_num in sorted_q_nums:
            row_data.append(q_map.get(q_num, "-"))
        row_data.extend([
            total_score,
            max_total,
            flagged_count,
            f"{avg_confidence:.2%}" if isinstance(avg_confidence, (int, float)) else str(avg_confidence),
            status,
        ])

        ws.append(row_data)

        # Style data row
        for col_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = thin_border
            cell.alignment = Alignment(
                horizontal="left" if col_idx <= 2 else "center",
                vertical="center"
            )
            # Highlight status column
            if col_idx == len(row_data):
                if status in status_fills:
                    cell.fill = status_fills[status]

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # Save to memory buffer
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"AutoGrade_Results_{exam_id}.xlsx"
    headers = {
        "Content-Disposition": f'attachment; filename="{filename}"'
    }

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )
